import json
import csv
import io
import logging
import sqlite3
from datetime import datetime, date, timedelta
from collections import defaultdict
from flask import Blueprint, request, jsonify, send_file, g
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

logger = logging.getLogger(__name__)

from utils import (
    get_db_connection, get_table_columns, parse_db_date, parse_db_datetime, cache_get, cache_set,
    vendor_has_feature, _run
)
from services.attendance_service import (
    calculate_daily_hours, calculate_expected_hours, calculate_arrival_status
)
from services.payroll_service import calculate_salary_breakdown, get_pending_advances
from services.auth_service import require_auth
from services.report_filter_service import (
    STANDARD_FILTERS, custom_value, face_matches, merge_filter_configuration, parse_json_list,
)
from middleware.validation import validate_request
from schemas import PayrollReportRequest

attendance_reports_bp = Blueprint('attendance_reports_bp', __name__)

# ---------------------------------------------------------------------------
# Helper: parse dynamic filters from request args (dynamic_<field>=value)
# ---------------------------------------------------------------------------
def _parse_dynamic_args():
    """Return dict of dynamic field filters from query params like dynamic_site=PlantA."""
    dynamic = {}
    for key, val in request.args.items():
        if key.startswith('dynamic_') and val:
            dynamic[key[len('dynamic_'):]] = val
    return dynamic


# ---------------------------------------------------------------------------
# /reports/filters  –  mirrors /attendance/filters but under the reports prefix
# ---------------------------------------------------------------------------
@attendance_reports_bp.route("/reports/filters", methods=["GET"])
@require_auth()
def get_report_filters():
    vendor_id = g.vendor_id

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    # ---- vendor meta ----
    c.execute("SELECT company_name, registration_config FROM vendors WHERE id = ?", (vendor_id,))
    vendor_row = c.fetchone()
    vendor_name = vendor_row['company_name'] if vendor_row else ''
    raw_reg = vendor_row['registration_config'] if vendor_row else None
    
    # Registration configuration and bulk-upload fields are the only authorities
    # for report filter visibility. The data itself supplies option values.
    hydrated_registration = []
    if raw_reg:
        try:
            config_data = json.loads(raw_reg) if isinstance(raw_reg, str) else raw_reg
            from services.config_utils import hydrate_registration_config
            hydrated_registration = hydrate_registration_config(vendor_id, config_data, conn=conn)
        except Exception:
            logger.exception("Unable to read report registration configuration for vendor %s", vendor_id)

    bulk_fields = []
    try:
        c.execute("SELECT fields FROM bulk_attendance_config WHERE vendor_id = ?", (vendor_id,))
        bulk_row = c.fetchone()
        if bulk_row:
            bulk_fields = parse_json_list(bulk_row['fields'])
    except Exception:
        logger.exception("Unable to read bulk report fields for vendor %s", vendor_id)
    visible_standard_filters, standard_filter_labels, enabled_dynamic_fields = merge_filter_configuration(
        hydrated_registration, bulk_fields,
    )

    # ---- collect active request-level filters for cascading ----
    req_standard = {key: request.args.get(key, '') for key in STANDARD_FILTERS}
    req_dyn   = _parse_dynamic_args()

    # ---- fetch all faces for this vendor ----
    c.execute(
        "SELECT department, designation, shift, phone, custom_data "
        "FROM faces WHERE vendor_id = ?",
        (vendor_id,)
    )
    faces_raw = c.fetchall()
    conn.close()

    faces = []
    for r in faces_raw:
        d = dict(r)
        try:
            d["custom"] = json.loads(d.get("custom_data") or "{}")
        except Exception:
            d["custom"] = {}
        faces.append(d)

    standard_options = {}
    for key in STANDARD_FILTERS:
        candidates = [face for face in faces if face_matches(face, req_standard, req_dyn, exclude_standard=key)]
        standard_options[key] = sorted({str(face.get(key)).strip() for face in candidates if face.get(key)}) if visible_standard_filters[key] else []

    # ---- build dynamic filter option lists ----
    dynamic_filters = {}
    for field in enabled_dynamic_fields:
        fk, fl = field['key'], field['label']
        candidates = [face for face in faces if face_matches(face, req_standard, req_dyn, exclude_dynamic=fk)]
        unique_values = set()
        for f in candidates:
            val = custom_value(f['custom'], fk)
            if val is not None and str(val).strip():
                unique_values.add(str(val).strip())
        options = sorted(list(unique_values))[:200]
        # if the config defines a fixed option list, always show all configured options
        if field.get('options'):
            options = [str(x) for x in field['options']]
        dynamic_filters[fk] = {'label': fl, 'options': options}

    return jsonify({
        'vendor_name': vendor_name,
        'departments': standard_options['department'],
        'designations': standard_options['designation'],
        'shifts': standard_options['shift'],
        'phones': standard_options['phone'],
        'visible_standard_filters': visible_standard_filters,
        'standard_filter_labels': standard_filter_labels,
        'dynamic_filters': dynamic_filters,
    })

@attendance_reports_bp.route("/reports/analytics", methods=["GET"])
@require_auth()
@validate_request(PayrollReportRequest)
def get_analytics(valid_data: PayrollReportRequest):
    vendor_id = g.vendor_id
    
    start_date = valid_data.start_date or (datetime.now() - timedelta(days=6)).strftime('%Y-%m-%d')
    end_date = valid_data.end_date or datetime.now().strftime('%Y-%m-%d')
        
    cache_key = f"vendor:{vendor_id}:analytics:{start_date}:{end_date}"
    cached = cache_get(cache_key)
    if cached: return jsonify(cached)

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    
    # 1. Fetch Timetable
    timetable = []
    c.execute("SELECT live_timetable FROM companies WHERE vendor_id = ?", (vendor_id,))
    company_row = c.fetchone()
    if company_row and company_row['live_timetable']:
        try: timetable = json.loads(company_row['live_timetable'])
        except (json.JSONDecodeError, ValueError): pass

    # 2. Fetch all faces, then separate faculty from students.
    #    Faculty upload creates face records too, so we must exclude them from
    #    the student attendance denominator (absent/present calculations).
    c.execute("SELECT id, name, department, designation FROM faces WHERE vendor_id = ?", (vendor_id,))
    all_faces = {row['id']: dict(row) for row in c.fetchall()}

    # Identify faculty person_ids via system_users
    c.execute("""
        SELECT person_id FROM system_users
        WHERE vendor_id = ? AND role = 'faculty' AND person_id IS NOT NULL
    """, (vendor_id,))
    faculty_ids = {row[0] for row in c.fetchall()}

    total_faculty = len(faculty_ids)
    total_all = len(all_faces)

    # Analytics (present/absent/late) runs on students only —
    # faculty face records skew absent_today if they never check in.
    persons = {pid: info for pid, info in all_faces.items() if pid not in faculty_ids}
    total_students = len(persons)
    
    # 3. Fetch Attendance
    try:
        s_dt = datetime.strptime(start_date, '%Y-%m-%d')
        e_dt = datetime.strptime(end_date, '%Y-%m-%d')
    except ValueError:
        conn.close(); return jsonify({"error": "Invalid date format"}), 400

    query = "SELECT * FROM attendance WHERE date(timestamp) BETWEEN ? AND ? AND vendor_id = ? ORDER BY timestamp ASC"
    c.execute(query, (start_date, end_date, vendor_id))
    rows = c.fetchall()
    conn.close()
    
    user_records = defaultdict(list)
    for row in rows:
        record = dict(row) if hasattr(row, 'keys') else dict(zip([d[0] for d in c.description], row))
        pid = record.get('person_id')
        if pid: user_records[pid].append(record)
        
    # Analysis
    daily_stats = defaultdict(lambda: {"present": 0, "late": 0, "undertime": 0, "overtime": 0})
    person_stats = {}
    
    total_days = (e_dt - s_dt).days + 1
    all_dates = [(s_dt + timedelta(days=x)).strftime('%Y-%m-%d') for x in range(total_days)]
    
    for pid, person_info in persons.items():
        records = user_records.get(pid, [])
        p_total_hours = 0
        p_present_days = 0
        p_late_count = 0
        
        # Group records by date for this person
        p_date_records = defaultdict(list)
        for r in records:
            d_str = parse_db_date(r['timestamp'])
            if d_str: p_date_records[d_str.strftime('%Y-%m-%d')].append(r)
            
        for d_str in all_dates:
            d_records = p_date_records.get(d_str, [])
            if not d_records: continue
            
            stats = calculate_daily_hours(d_records, timetable, date_str=d_str)
            p_total_hours += stats['total_hours']
            p_present_days += 1
            
            # Check if late in any session of the day
            is_late_day = any(r.get('is_late') == 1 for r in d_records)
            if is_late_day: p_late_count += 1
            
            # Update daily global stats
            daily_stats[d_str]["present"] += 1
            if is_late_day: daily_stats[d_str]["late"] += 1
            
        person_stats[pid] = {
            "name": person_info['name'],
            "department": person_info['department'],
            "total_hours": round(p_total_hours, 2),
            "present_days": p_present_days,
            "late_count": p_late_count,
            "attendance_rate": round((p_present_days / total_days) * 100, 1) if total_days > 0 else 0
        }
    
    total_persons = len(persons)
    today_str = datetime.now().strftime('%Y-%m-%d')
    present_today = daily_stats[today_str]["present"]
    late_today = daily_stats[today_str]["late"]
    absent_today = max(0, total_students - present_today)
    on_time_today = max(0, present_today - late_today)

    pie_data = [
        {"name": "On Time", "value": on_time_today},
        {"name": "Late", "value": late_today},
        {"name": "Absent", "value": absent_today}
    ]

    bar_data = []
    for d_str in all_dates:
        p = daily_stats[d_str]["present"]
        l = daily_stats[d_str]["late"]
        a = max(0, total_students - p)
        try:
            d_name = datetime.strptime(d_str, '%Y-%m-%d').strftime('%a')
        except (ValueError, TypeError):
            d_name = d_str
        bar_data.append({
            "name": d_name, "date": d_str, "present": p, "absent": a, "late": l, "total": total_students
        })

    result = {
        "daily_stats": daily_stats,
        "person_stats": person_stats,
        "bar_data": bar_data,
        "pie_data": pie_data,
        "dept_data": [],
        "summary": {
            "total_users": total_all,
            "total_students": total_students,
            "total_faculty": total_faculty,
            "present_today": present_today,
            "absent_today": absent_today,   # based on students only
            "late_today": late_today,
            "avg_attendance": round(sum(p['attendance_rate'] for p in person_stats.values()) / len(person_stats), 1) if person_stats else 0
        }
    }
    cache_set(cache_key, result, 300)
    return jsonify(result)

@attendance_reports_bp.route("/reports/export", methods=["GET"])
@require_auth()
@validate_request(PayrollReportRequest)
def export_attendance(valid_data: PayrollReportRequest):
    vendor_id = g.vendor_id

    start_date = valid_data.start_date
    end_date   = valid_data.end_date

    # --- active filters ---
    f_dept  = request.args.get('department', '').strip()
    f_desig = request.args.get('designation', '').strip()
    f_shift = request.args.get('shift', '').strip()
    f_phone = request.args.get('phone', '').strip()
    dyn_filters = _parse_dynamic_args()  # {field_key: value, ...}

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    # Build query with optional standard-filter JOINs on faces
    query = """
        SELECT a.name, a.timestamp, a.status, a.activity, a.is_late,
               f.department, f.designation, f.shift, f.phone, f.custom_data
        FROM attendance a
        LEFT JOIN faces f ON a.person_id = f.id
        WHERE date(a.timestamp) BETWEEN ? AND ? AND a.vendor_id = ?
    """
    params = [start_date, end_date, vendor_id]

    if f_dept:
        query += " AND f.department = ?"; params.append(f_dept)
    if f_desig:
        query += " AND f.designation = ?"; params.append(f_desig)
    if f_shift:
        query += " AND f.shift = ?"; params.append(f_shift)
    if f_phone:
        query += " AND f.phone = ?"; params.append(f_phone)

    query += " ORDER BY a.timestamp ASC"
    c.execute(query, params)
    db_rows = c.fetchall()
    conn.close()

    # Extract unique custom-data keys and apply dynamic filters
    custom_keys  = []
    parsed_rows  = []
    for row in db_rows:
        row_dict  = dict(row)
        cdata_str = row_dict.pop('custom_data', None)
        cdata     = {}
        if cdata_str:
            try:
                cdata = json.loads(cdata_str)
                if isinstance(cdata, dict):
                    for k in cdata.keys():
                        if k not in custom_keys:
                            custom_keys.append(k)
            except Exception:
                pass
        row_dict['parsed_custom'] = cdata

        # Apply dynamic (custom_data) filters – Python-side after parse
        match = True
        for dk, dv in dyn_filters.items():
            if str(custom_value(cdata, dk) or '').strip() != dv:
                match = False
                break
        if not match:
            continue

        parsed_rows.append(row_dict)

    # Build CSV
    base_headers     = ["Name", "Date", "Time", "Status", "Activity", "Is Late",
                        "Department", "Designation", "Shift", "Phone"]
    formatted_custom = [k.replace('_', ' ').title() for k in custom_keys]
    headers          = base_headers + formatted_custom

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)

    for row in parsed_rows:
        d     = parse_db_datetime(row['timestamp'])
        d_str = d.strftime('%Y-%m-%d') if d else ""
        t_str = d.strftime('%H:%M:%S') if d else ""

        base_record = [
            row['name'], d_str, t_str, row['status'], row['activity'],
            "Yes" if row['is_late'] else "No",
            row.get('department', ''), row.get('designation', ''),
            row.get('shift', ''), row.get('phone', '')
        ]
        custom_record = [row['parsed_custom'].get(k, '') for k in custom_keys]
        writer.writerow(base_record + custom_record)

    return output.getvalue(), 200, {
        "Content-Type": "text/csv",
        "Content-Disposition": f"attachment; filename=attendance_{start_date}_{end_date}.csv"
    }

@attendance_reports_bp.route("/reports/payroll", methods=["GET"])
@require_auth()
@validate_request(PayrollReportRequest)
def get_payroll_report(valid_data: PayrollReportRequest):
    vendor_id = g.vendor_id

    start_date = valid_data.start_date
    end_date = valid_data.end_date
        
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    
    # 1. Fetch Timetable and Working Hours
    c.execute("SELECT live_timetable, working_hours FROM companies WHERE vendor_id = ?", (vendor_id,))
    company_row = c.fetchone()

    timetable = []
    company_working_hours = 8.0 # Default
    if company_row:
        if company_row['live_timetable']:
            try: timetable = json.loads(company_row['live_timetable'])
            except (json.JSONDecodeError, ValueError): pass
        try:
            if company_row['working_hours']: company_working_hours = float(company_row['working_hours'])
        except (ValueError, TypeError): pass

    late_enabled = vendor_has_feature(vendor_id, "late_mark")
    
    # 1. Incoming global percentages (Query params override stored settings)
    # Fetch stored settings first
    settings_keys = [
        'global_late_allowance', 
        'global_late_deduction',
        'global_pf_percentage',
        'global_esi_percentage',
        'global_gratuity_percentage',
        'global_gratuity_threshold_years',
        'global_timezone_offset'
    ]
    try:
        # DB Factory handles the ? placeholder correctly for both PG and SQLite
        vendor_setting_keys = [f"{key}_vendor_{vendor_id}" for key in settings_keys]
        placeholders = ', '.join(['?'] * len(vendor_setting_keys))
        c.execute(f"SELECT key, value FROM system_settings WHERE key IN ({placeholders})", vendor_setting_keys)
        stored_settings = {
            row['key'].removesuffix(f"_vendor_{vendor_id}"): row['value']
            for row in c.fetchall()
        }
    except Exception:
        logger.warning("Failed to load system settings", exc_info=True)
        stored_settings = {}

    pf_pct = request.args.get('pf_percentage')
    if pf_pct is None: pf_pct = stored_settings.get('global_pf_percentage', '12.0')
    
    esi_pct = request.args.get('esi_percentage')
    if esi_pct is None: esi_pct = stored_settings.get('global_esi_percentage', '0.75')
    
    gratuity_pct = request.args.get('gratuity_percentage')
    if gratuity_pct is None: gratuity_pct = stored_settings.get('global_gratuity_percentage', '4.81')
    
    grat_years_threshold = stored_settings.get('global_gratuity_threshold_years')
    if grat_years_threshold is None: grat_years_threshold = 5
    else: grat_years_threshold = int(grat_years_threshold)

    try:
        pf_pct = float(pf_pct)
        esi_pct = float(esi_pct)
        gratuity_pct = float(gratuity_pct)
    except (ValueError, TypeError):
        pf_pct = 12.0
        esi_pct = 0.75
        gratuity_pct = 4.81

    # 2. Fetch Persons
    try:
        fcols = get_table_columns(conn, "faces")
        payroll_fields = ["basic_salary", "hra", "conveyance", "special_allowance", "pf_enabled", "esi_enabled", "gratuity_enabled", "professional_tax", "joining_date"]
        extra = [col for col in (["late_allowance_days", "late_deduction_amount"] + payroll_fields) if col in fcols]
        extra_select = ", " + ", ".join(extra) if extra else ""
        c.execute(f"SELECT id, name, daily_wage, department, designation, face_image, phone, display_id, custom_data{extra_select} FROM faces WHERE vendor_id = ? ORDER BY id ASC", (vendor_id,))
    except Exception:
        logger.debug("Extended faces columns unavailable, falling back to base columns", exc_info=True)
        c.execute("SELECT id, name, daily_wage, department, designation, face_image, phone FROM faces WHERE vendor_id = ? ORDER BY id ASC", (vendor_id,))
    
    persons = {row['id']: dict(row) for row in c.fetchall()}
    
    global_allowance = int(stored_settings.get('global_late_allowance', 7))
    global_deduction = float(stored_settings.get('global_late_deduction', 0.0))
    
    # 4. Fetch Attendance
    try:
        s_dt = datetime.strptime(start_date, '%Y-%m-%d') - timedelta(days=1)
        e_dt = datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1)
    except ValueError:
        conn.close(); return jsonify({"error": "Invalid date format"}), 400

    query = "SELECT * FROM attendance WHERE date(timestamp) BETWEEN ? AND ? AND vendor_id = ? ORDER BY timestamp ASC"
    c.execute(query, (s_dt.strftime('%Y-%m-%d'), e_dt.strftime('%Y-%m-%d'), vendor_id))
    rows = c.fetchall()
    
    user_records = defaultdict(list)
    for row in rows:
        pid = row.get('person_id')
        if pid: user_records[pid].append(dict(row))

    # Fetch list of advances for marking status later if needed (but here we just read)
    # The connection is still open.
        
    payroll_data = []
    # today_str = datetime.now().strftime('%Y-%m-%d') # Removed unused
    start_dt_req = datetime.strptime(start_date, '%Y-%m-%d').date()
    end_dt_req = datetime.strptime(end_date, '%Y-%m-%d').date()

    for pid, person_info in persons.items():
        total_hours = 0
        present_dates = set()
        late_dates = set()
        
        records = user_records.get(pid, [])
        stats = calculate_daily_hours(records, timetable)
        
        for r in records:
            if r.get('is_late') == 1:
                r_dt = parse_db_date(r['timestamp'])
                if r_dt and start_dt_req <= r_dt <= end_dt_req:
                    late_dates.add(r_dt)
        
        late_marks_count = len(late_dates) if late_enabled else 0

        for session in stats.get('sessions', []):
            if session.get('start_ts'):
                try:
                    sess_start = datetime.fromisoformat(session['start_ts'])
                    sess_date = sess_start.date()
                    if start_dt_req <= sess_date <= end_dt_req and session.get('is_payable', False):
                        total_hours += (session.get('duration_mins', 0) / 60.0)
                        present_dates.add(sess_date)
                except (ValueError, TypeError): continue

        daily_wage = person_info.get('daily_wage') or 0
        hourly_rate = daily_wage / company_working_hours if daily_wage and company_working_hours > 0 else 0
        base_cost = round(total_hours * hourly_rate, 2)
        
        allowance = person_info.get('late_allowance_days') if person_info.get('late_allowance_days') is not None else global_allowance
        deduction_amt = person_info.get('late_deduction_amount') if person_info.get('late_deduction_amount') is not None else global_deduction
        
        deductable_lates = max(0, late_marks_count - allowance)
        total_deduction = round(deductable_lates * deduction_amt, 2)
        
        # Get Month for advances (using end_date if it covers the month)
        target_month = end_dt_req.strftime('%Y-%m')
        
        # Determine Tenure (Years) for Gratuity Threshold
        tenure_years = 0
        if person_info.get('joining_date'):
            try:
                # Handling both string and date/datetime objects for robustness
                j_raw = person_info['joining_date']
                if isinstance(j_raw, str):
                    jdate = datetime.strptime(j_raw, '%Y-%m-%d').date()
                elif hasattr(j_raw, 'date'):
                    jdate = j_raw.date()
                else:
                    jdate = j_raw
                
                today = date.today()
                tenure_years = (today - jdate).days / 365.25
            except (ValueError, TypeError, AttributeError):
                pass
        
        # Calculate Breakdown using the service
        person_info_with_tenure = person_info.copy()
        person_info_with_tenure['tenure_years'] = tenure_years
        
        breakdown = calculate_salary_breakdown(
            base_cost, 
            person_info_with_tenure, 
            pf_percent=pf_pct, 
            esi_percent=esi_pct, 
            gratuity_percent=gratuity_pct,
            gratuity_threshold_years=grat_years_threshold
        )
        
        # Fetch and process advances
        advances = get_pending_advances(conn, pid, target_month)
        advance_total = sum(adv[1] for adv in advances)
        
        final_payout = round(breakdown['net_before_advances'] - total_deduction - advance_total, 2)
        
        payroll_data.append({
            "person_id": pid,
            "display_id": person_info.get('display_id') or pid,
            "name": person_info.get('name'),
            "department": person_info.get('department'),
            "designation": person_info.get('designation'),
            "face_image": person_info.get('face_image'),
            "phone": person_info.get('phone'),
            "daily_wage": daily_wage,
            "total_hours": round(total_hours, 2),
            "days_present": len(present_dates),
            "base_cost": base_cost,
            "late_marks_count": late_marks_count,
            "late_deduction": total_deduction,
            "advance_deduction": advance_total,
            "pf_enabled": person_info.get('pf_enabled', 0),
            "esi_enabled": person_info.get('esi_enabled', 0),
            "gratuity_enabled": person_info.get('gratuity_enabled', 0),
            "professional_tax": person_info.get('professional_tax', 0),
            "late_allowance_days": person_info.get('late_allowance_days'),
            "late_deduction_amount": person_info.get('late_deduction_amount'),
            "breakdown": breakdown,
            "final_payout": final_payout,
            "custom_data": json.loads(person_info.get('custom_data') or '{}') if isinstance(person_info.get('custom_data'), str) else person_info.get('custom_data')
        })
    
    conn.close()
    return jsonify({
        "payroll": payroll_data,
        "global_settings": {
            "allowance": global_allowance, 
            "deduction": global_deduction, 
            "pf_percentage": pf_pct, 
            "esi_percentage": esi_pct, 
            "gratuity_percentage": gratuity_pct,
            "gratuity_threshold_years": grat_years_threshold
        }
    })

@attendance_reports_bp.route("/reports/payroll/export-daily", methods=["GET"])
@require_auth()
def export_payroll_daily():
    vendor_id  = g.vendor_id
    start_date = request.args.get("start_date")
    end_date   = request.args.get("end_date")
    if not start_date or not end_date:
        return jsonify({"error": "start_date and end_date required"}), 400

    # active standard filters
    f_dept  = request.args.get('department', '').strip()
    f_desig = request.args.get('designation', '').strip()
    f_shift = request.args.get('shift', '').strip()
    f_phone = request.args.get('phone', '').strip()
    dyn_filters = _parse_dynamic_args()

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    try:
        c.execute("SELECT working_hours FROM companies WHERE vendor_id = ?", (vendor_id,))
        v_row = c.fetchone()
        company_working_hours = float(v_row['working_hours']) if v_row and v_row['working_hours'] else 8.0

        # Add a 1-day buffer on each side to capture timezone-shifted timestamps
        start_dt_req = datetime.strptime(start_date, '%Y-%m-%d').date()
        end_dt_req   = datetime.strptime(end_date, '%Y-%m-%d').date()
        fetch_start  = (start_dt_req - timedelta(days=1)).isoformat()
        fetch_end    = (end_dt_req   + timedelta(days=1)).isoformat()

        c.execute(
            "SELECT * FROM attendance WHERE vendor_id = ? AND date(timestamp) BETWEEN ? AND ?",
            (vendor_id, fetch_start, fetch_end)
        )
        rows = c.fetchall()
        user_records = defaultdict(list)
        for row in rows:
            d = dict(row)
            if d.get('person_id'): user_records[d['person_id']].append(d)

        # Build faces query with standard filters
        fcols = get_table_columns(conn, "faces")
        payroll_fields = ["basic_salary", "hra", "conveyance", "special_allowance", "pf_enabled", "esi_enabled", "professional_tax"]
        extra = [col for col in (["late_allowance_days", "late_deduction_amount"] + payroll_fields) if col in fcols]
        extra_select = ", " + ", ".join(extra) if extra else ""
        
        faces_query  = f"SELECT id, name, daily_wage, shift, department, designation, phone, custom_data, display_id{extra_select} FROM faces WHERE vendor_id = ?"
        faces_params = [vendor_id]
        if f_dept:  faces_query += " AND department = ?";  faces_params.append(f_dept)
        if f_desig: faces_query += " AND designation = ?"; faces_params.append(f_desig)
        if f_shift: faces_query += " AND shift = ?" ;      faces_params.append(f_shift)
        if f_phone: faces_query += " AND phone = ?" ;      faces_params.append(f_phone)

        c.execute(faces_query, faces_params)
        all_faces = c.fetchall()

        # Apply dynamic (custom_data) filter if any
        persons = {}
        for r in all_faces:
            rd = dict(r)
            if dyn_filters:
                try:
                    cdata = json.loads(rd.get('custom_data') or '{}')
                except Exception:
                    cdata = {}
                if any(str(custom_value(cdata, dk) or '').strip() != dv for dk, dv in dyn_filters.items()):
                    continue
            persons[rd['id']] = rd

        setting_suffix = f'_vendor_{vendor_id}'
        c.execute("SELECT key, value FROM system_settings WHERE key IN (?, ?)",
                  (f'global_late_allowance{setting_suffix}', f'global_late_deduction{setting_suffix}'))
        s_map = {r['key'].removesuffix(setting_suffix): r['value'] for r in c.fetchall()}
        global_allowance = int(s_map.get('global_late_allowance', 7))
        global_deduction = float(s_map.get('global_late_deduction', 0.0))

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Employee ID", "Name", "Date", "Daily Wage Rate", "Working Hours", "Base Wage", "Late Penalty", "PF", "ESI", "PT", "Net Wage"])

        for pid, records in user_records.items():
            if pid not in persons: continue
            info = persons[pid]
            # Simple fallback for timetable if needed (would ideally use company timetable)
            timetable = []
            try:
                tjson = info.get('shift') or ""
                if tjson: timetable = json.loads(tjson) if isinstance(tjson, str) else (tjson or [])
            except (json.JSONDecodeError, ValueError): pass
            
            stats = calculate_daily_hours(records, timetable)
            per_day = defaultdict(float)
            for s in stats.get('sessions', []):
                try:
                    sd = datetime.fromisoformat(s['start_ts']).date()
                    if start_dt_req <= sd <= end_dt_req and s.get('is_payable', False):
                        per_day[sd] += (s.get('duration_mins', 0) or 0)
                except (ValueError, TypeError): continue

            late_dates = sorted(set(parse_db_date(r['timestamp']) for r in records if r.get('is_late') == 1 and parse_db_date(r['timestamp']) and start_dt_req <= parse_db_date(r['timestamp']) <= end_dt_req))
            allowance = info.get('late_allowance_days') if info.get('late_allowance_days') is not None else global_allowance
            deduction_amt = info.get('late_deduction_amount') if info.get('late_deduction_amount') is not None else global_deduction

            for d, mins in sorted(per_day.items()):
                lm = 1 if d in late_dates else 0
                idx = late_dates.index(d) + 1 if d in late_dates else 0
                deduct = deduction_amt if (lm == 1 and idx > allowance) else 0.0
                daily_wage = float(info.get('daily_wage') or 0.0)
                hourly_rate = daily_wage / company_working_hours if daily_wage and company_working_hours > 0 else 0
                base_wage_for_day = round((mins / 60.0) * hourly_rate, 2)
                
                # Breakdown for this day's portion
                breakdown = calculate_salary_breakdown(base_wage_for_day, info)
                pf = breakdown['deductions']['pf']
                esi = breakdown['deductions']['esi']
                pt = breakdown['deductions']['pt'] / 26.0 # Arbitrary per day PT if monthly
                
                net_wage = round(breakdown['net_before_advances'] - deduct, 2)
                d_id = info.get('display_id') or pid
                writer.writerow([d_id, info.get('name'), d.isoformat(), daily_wage, round(mins/60.0,2), base_wage_for_day, deduct, pf, esi, round(pt, 2), net_wage])
        
        return output.getvalue(), 200, {"Content-Type": "text/csv"}
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        conn.close()

@attendance_reports_bp.route("/reports/payroll/export-excel", methods=["GET"])
@require_auth()
def export_payroll_excel():
    vendor_id = g.vendor_id
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    if not start_date or not end_date:
        return jsonify({"error": "start_date and end_date required"}), 400

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    try:
        c.execute("SELECT working_hours FROM companies WHERE vendor_id = ?", (vendor_id,))
        v_row = c.fetchone()
        company_working_hours = float(v_row['working_hours']) if v_row and v_row['working_hours'] else 8.0

        # Add a 1-day buffer on each side to capture timezone-shifted timestamps
        start_dt_req = datetime.strptime(start_date, '%Y-%m-%d').date()
        end_dt_req   = datetime.strptime(end_date, '%Y-%m-%d').date()
        fetch_start  = (start_dt_req - timedelta(days=1)).isoformat()
        fetch_end    = (end_dt_req   + timedelta(days=1)).isoformat()

        c.execute(
            "SELECT * FROM attendance WHERE vendor_id = ? AND date(timestamp) BETWEEN ? AND ?",
            (vendor_id, fetch_start, fetch_end)
        )
        rows = c.fetchall()
        user_records = defaultdict(list)
        for row in rows:
            d = dict(row)
            if d.get('person_id'): user_records[d['person_id']].append(d)

        fcols = get_table_columns(conn, "faces")
        payroll_fields = ["basic_salary", "hra", "conveyance", "special_allowance", "pf_enabled", "esi_enabled", "professional_tax"]
        extra = [col for col in (["late_allowance_days", "late_deduction_amount"] + payroll_fields) if col in fcols]
        extra_select = ", " + ", ".join(extra) if extra else ""
        
        c.execute(f"SELECT id, name, daily_wage, display_id{extra_select} FROM faces WHERE vendor_id = ?", (vendor_id,))
        persons = {r['id']: dict(r) for r in c.fetchall()}

        setting_suffix = f'_vendor_{vendor_id}'
        c.execute("SELECT key, value FROM system_settings WHERE key IN (?, ?)",
                  (f'global_late_allowance{setting_suffix}', f'global_late_deduction{setting_suffix}'))
        s_map = {r['key'].removesuffix(setting_suffix): r['value'] for r in c.fetchall()}
        global_allowance = int(s_map.get('global_late_allowance', 7))
        global_deduction = float(s_map.get('global_late_deduction', 0.0))

        wb = Workbook()
        ws = wb.active
        ws.title = "Daily Payroll"
        headers = ["Employee ID", "Name", "Date", "Daily Wage Rate", "Working Hours", "Base Wage", "Late Penalty", "PF", "ESI", "PT", "Net Wage"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        start_dt_req = datetime.fromisoformat(start_date).date()
        end_dt_req = datetime.fromisoformat(end_date).date()

        for pid, records in user_records.items():
            if pid not in persons: continue
            info = persons[pid]
            stats = calculate_daily_hours(records)
            per_day = defaultdict(float)
            for s in stats.get('sessions', []):
                try:
                    sd = datetime.fromisoformat(s['start_ts']).date()
                    if start_dt_req <= sd <= end_dt_req and s.get('is_payable', False):
                        per_day[sd] += (s.get('duration_mins', 0) or 0)
                except (ValueError, TypeError): continue

            late_dates = sorted(set(parse_db_date(r['timestamp']) for r in records if r.get('is_late') == 1 and parse_db_date(r['timestamp']) and start_dt_req <= parse_db_date(r['timestamp']) <= end_dt_req))
            allowance = info.get('late_allowance_days') if info.get('late_allowance_days') is not None else global_allowance
            deduction_amt = info.get('late_deduction_amount') if info.get('late_deduction_amount') is not None else global_deduction

            for d, mins in sorted(per_day.items()):
                lm = 1 if d in late_dates else 0
                idx = late_dates.index(d) + 1 if d in late_dates else 0
                deduct = deduction_amt if (lm == 1 and idx > allowance) else 0.0
                daily_wage = float(info.get('daily_wage') or 0.0)
                hourly_rate = daily_wage / company_working_hours if daily_wage and company_working_hours > 0 else 0
                base_wage_for_day = round((mins / 60.0) * hourly_rate, 2)
                
                # Breakdown
                breakdown = calculate_salary_breakdown(base_wage_for_day, info)
                pf = breakdown['deductions']['pf']
                esi = breakdown['deductions']['esi']
                pt = breakdown['deductions']['pt'] / 26.0
                
                net_wage = round(breakdown['net_before_advances'] - deduct, 2)
                d_id = info.get('display_id') or pid
                ws.append([d_id, info.get('name'), d.isoformat(), daily_wage, round(mins/60.0, 2), base_wage_for_day, deduct, pf, esi, round(pt, 2), net_wage])

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"payroll_daily_{start_date}_{end_date}.xlsx")
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        conn.close()

@attendance_reports_bp.route("/reports/payroll/import", methods=["POST"])
@require_auth()
def import_payroll_adjustments():
    vendor_id = g.vendor_id
    if 'file' not in request.files:
        return jsonify({"error":"CSV file required"}), 400
    f = request.files['file']
    try:
        content = f.read().decode('utf-8')
        reader = csv.DictReader(io.StringIO(content))
        conn = get_db_connection()
        c = conn.cursor()
        updated = 0
        for row in reader:
            pid = row.get('person_id')
            if not pid: continue
            vals = {}
            try:
                if row.get('late_allowance_days'): vals['late_allowance_days'] = int(row['late_allowance_days'])
                if row.get('late_deduction_amount'): vals['late_deduction_amount'] = float(row['late_deduction_amount'])
            except (ValueError, TypeError): pass
            if vals:
                c.execute("UPDATE faces SET late_allowance_days = ?, late_deduction_amount = ? WHERE vendor_id = ? AND id = ?", (vals.get('late_allowance_days'), vals.get('late_deduction_amount'), vendor_id, pid))
                updated += 1
        conn.commit(); conn.close()
        return jsonify({"success": True, "updated": updated})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
